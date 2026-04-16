"""
generate_photosheet.py
Reads a photo log (Excel or CSV) and a dict of {filename: bytes},
and produces a .docx file with two photo blocks per page.
"""

import io
import csv
import os
from typing import Dict, List, Optional, Tuple

from PIL import Image, ExifTags
from docx import Document
from docx.shared import Inches, Pt

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import pandas as pd


# ------------------------------------------------------------------
# Page geometry (Letter, portrait, 1" margins all sides)
# ------------------------------------------------------------------
PAGE_WIDTH_IN  = 8.5
PAGE_HEIGHT_IN = 11.0
MARGIN_IN      = 1.0

PRINTABLE_W = PAGE_WIDTH_IN  - 2 * MARGIN_IN   # 6.5"
PRINTABLE_H = PAGE_HEIGHT_IN - 2 * MARGIN_IN   # 9.0"

# Each photo block gets half the printable height minus a small gap
BLOCK_H     = (PRINTABLE_H / 2) - 0.15         # ~4.35"
# Reserve space for caption inside the block
CAPTION_H   = 0.35                              # inches
PHOTO_MAX_H = BLOCK_H - CAPTION_H              # ~4.0"
PHOTO_MAX_W = PRINTABLE_W                       # 6.5"


# ------------------------------------------------------------------
# Helpers
# ------------------------------------------------------------------

def _correct_orientation(img: Image.Image) -> Image.Image:
    """Apply EXIF orientation so the image appears upright."""
    try:
        exif = img._getexif()
        if exif is None:
            return img
        orient_tag = next(
            (k for k, v in ExifTags.TAGS.items() if v == 'Orientation'), None
        )
        if orient_tag is None or orient_tag not in exif:
            return img
        orientation = exif[orient_tag]
        rotations = {
            3: 180,
            6: 270,
            8: 90,
        }
        flips = {
            2: Image.FLIP_LEFT_RIGHT,
            4: Image.FLIP_TOP_BOTTOM,
            5: None,
            7: None,
        }
        if orientation in rotations:
            img = img.rotate(rotations[orientation], expand=True)
        elif orientation == 5:
            img = img.rotate(270, expand=True).transpose(Image.FLIP_LEFT_RIGHT)
        elif orientation == 7:
            img = img.rotate(90, expand=True).transpose(Image.FLIP_LEFT_RIGHT)
        elif orientation in flips and flips[orientation]:
            img = img.transpose(flips[orientation])
    except Exception:
        pass
    return img


def _fit_dimensions(img_w_px: int, img_h_px: int, img_dpi: float) -> Tuple[float, float]:
    """
    Return (width_in, height_in) scaled to fit within PHOTO_MAX_W x PHOTO_MAX_H,
    preserving aspect ratio.
    """
    w_in = img_w_px / img_dpi
    h_in = img_h_px / img_dpi

    scale = min(PHOTO_MAX_W / w_in, PHOTO_MAX_H / h_in, 1.0)
    return w_in * scale, h_in * scale


def _prepare_image(img_bytes: bytes) -> Tuple[io.BytesIO, float, float]:
    """
    Open image bytes, correct orientation, convert to RGB JPEG, and return
    (BytesIO, width_inches, height_inches).
    """
    img = Image.open(io.BytesIO(img_bytes))
    img = _correct_orientation(img)

    dpi_info = img.info.get('dpi', (96, 96))
    try:
        dpi = float(dpi_info[0]) if dpi_info[0] > 0 else 96.0
    except Exception:
        dpi = 96.0

    w_in, h_in = _fit_dimensions(img.width, img.height, dpi)

    # Convert to RGB (handles RGBA, palette, etc.)
    if img.mode not in ('RGB',):
        img = img.convert('RGB')

    buf = io.BytesIO()
    img.save(buf, format='JPEG', quality=90, dpi=(dpi, dpi))
    buf.seek(0)
    return buf, w_in, h_in


# ------------------------------------------------------------------
# Log parsing
# ------------------------------------------------------------------

INCLUDE_YES = {'yes', 'true', '1', 'y', 'x'}

def _parse_log(log_bytes: bytes, log_filename: str) -> List[dict]:
    """
    Returns list of dicts with keys: filename (str), caption (str), include (bool).
    Columns are auto-detected (case-insensitive).
    """
    if log_filename.endswith('.csv'):
        text = log_bytes.decode('utf-8-sig', errors='replace')
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
    else:
        df = pd.read_excel(io.BytesIO(log_bytes), dtype=str)
        df = df.fillna('')
        rows = df.to_dict(orient='records')

    # Normalise column names for detection
    def find_col(row_keys, candidates):
        for k in row_keys:
            if k.strip().lower() in candidates:
                return k
        return None

    if not rows:
        return []

    keys = list(rows[0].keys())

    filename_col = find_col(keys, {
        'filename', 'file name', 'file', 'photo', 'photo name',
        'image', 'image name', 'name'
    })
    caption_col = find_col(keys, {
        'caption', 'description', 'desc', 'text', 'label', 'note', 'notes'
    })
    include_col = find_col(keys, {
        'include', 'use', 'selected', 'select', 'flag', 'print'
    })

    if filename_col is None:
        # Fall back: first column
        filename_col = keys[0]
    if caption_col is None:
        # Fall back: second column if exists
        caption_col = keys[1] if len(keys) > 1 else keys[0]

    result = []
    for row in rows:
        fname = str(row.get(filename_col, '') or '').strip()
        if not fname:
            continue
        caption = str(row.get(caption_col, '') or '').strip()

        if include_col:
            raw = str(row.get(include_col, '') or '').strip().lower()
            include = raw in INCLUDE_YES
        else:
            include = True  # include all when no flag column

        if include:
            result.append({'filename': fname, 'caption': caption})

    return result


# ------------------------------------------------------------------
# Document assembly
# ------------------------------------------------------------------

def _set_page_margins(doc: Document):
    section = doc.sections[0]
    section.page_width  = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.left_margin   = Inches(MARGIN_IN)
    section.right_margin  = Inches(MARGIN_IN)
    section.top_margin    = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)


def _para_space(doc: Document, space_before_pt: float = 0, space_after_pt: float = 0):
    """Add an empty paragraph with controlled spacing (used for vertical padding)."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before_pt)
    pf.space_after  = Pt(space_after_pt)
    pf.line_spacing = Pt(1)
    return p


def _add_photo_block(doc: Document, img_buf: Optional[io.BytesIO],
                     w_in: float, h_in: float,
                     caption: str, filename: str, is_missing: bool):
    """
    Add one photo block: image (or placeholder) + caption, both centered.
    """
    if is_missing:
        # Placeholder paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after  = Pt(0)
        run = p.add_run(f'MISSING PHOTO: {filename}')
        run.bold = True
        run.font.size = Pt(12)
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after  = Pt(0)
        run = p.add_run()
        run.add_picture(img_buf, width=Inches(w_in), height=Inches(h_in))

    # Caption paragraph
    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cpf = cap_p.paragraph_format
    cpf.space_before = Pt(4)
    cpf.space_after  = Pt(0)
    cap_run = cap_p.add_run(caption)
    cap_run.font.size = Pt(10)


def _add_page_break(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    run = p.add_run()
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    run._r.append(br)


# ------------------------------------------------------------------
# Public entry point
# ------------------------------------------------------------------

def generate_photosheet(log_bytes: bytes, log_filename: str,
                        photo_map: Dict[str, bytes]) -> bytes:
    """
    log_bytes    : raw bytes of the Excel/CSV log
    log_filename : original filename (used to detect .csv vs .xlsx)
    photo_map    : {lowercase_basename: bytes}

    Returns bytes of the generated .docx.
    """
    entries = _parse_log(log_bytes, log_filename)
    if not entries:
        raise ValueError('No rows found in photo log (or no rows marked for inclusion).')

    doc = Document()
    _set_page_margins(doc)

    # Remove default empty paragraph Word adds
    for p in doc.paragraphs:
        p._element.getparent().remove(p._element)

    # Process entries in pairs
    for i in range(0, len(entries), 2):
        pair = entries[i:i+2]

        for j, entry in enumerate(pair):
            fname_orig = entry['filename']
            fname_key  = fname_orig.lower()
            caption    = entry['caption']

            img_data = photo_map.get(fname_key)
            if img_data is None:
                _add_photo_block(doc, None, 0, 0, caption, fname_orig, is_missing=True)
            else:
                try:
                    img_buf, w_in, h_in = _prepare_image(img_data)
                    _add_photo_block(doc, img_buf, w_in, h_in, caption, fname_orig, is_missing=False)
                except Exception:
                    _add_photo_block(doc, None, 0, 0, caption, fname_orig, is_missing=True)

            # Between the two blocks on the same page: small spacer
            if j == 0 and len(pair) == 2:
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(8)
                sp.paragraph_format.space_after  = Pt(8)

        # Page break after each pair (but not after the very last pair)
        if i + 2 < len(entries):
            _add_page_break(doc)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
